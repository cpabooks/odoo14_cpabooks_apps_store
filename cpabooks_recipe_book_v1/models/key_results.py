# -*- coding: utf-8 -*-
from odoo import api, fields, models


class RecipeKeyResults(models.Model):
    _inherit = 'cpabooks.recipe'

    @api.model
    def get_key_results(self, date_from=None, date_to=None):
        labels = dict(self._fields['category'].selection)
        labels['none'] = 'No recipe'
        labels['unmatched'] = 'Closing with no recipe'
        today = fields.Date.context_today(self)
        date_from = fields.Date.to_date(date_from) if date_from else today.replace(day=1)
        date_to = fields.Date.to_date(date_to) if date_to else today
        selected = self.env.companies
        company_ids = selected.ids
        companies = {}
        seen_recipes = set()

        def money(value):
            return '{:,.2f}'.format(value or 0.0)

        def qty(value):
            return '{:,.2f}'.format(value or 0.0)

        def qty_show(value):
            if not value:
                return ''
            return qty(value)

        def percent(cost, sales):
            if not sales:
                return 0.0
            return (cost / sales) * 100.0

        def gap_text(sales, cost):
            if not sales and cost:
                return 'Sold amount missing'
            if sales and not cost:
                return 'Cost missing'
            if not sales and not cost:
                return 'No sales and no cost'
            return ''

        def company_bucket(company):
            row = companies.get(company.id)
            if row:
                return row
            row = {
                'id': company.id,
                'name': company.name,
                'sales': 0.0,
                'cost': 0.0,
                'sold_qty': 0.0,
                'theoretical_qty': 0.0,
                'purchase_qty': 0.0,
                'daily_qty': 0.0,
                'categories': {},
            }
            companies[company.id] = row
            return row

        def category_bucket(company_row, key):
            row = company_row['categories'].get(key)
            if row:
                return row
            row = {
                'key': key,
                'name': labels.get(key, key),
                'sold_qty': 0.0,
                'sales': 0.0,
                'cost': 0.0,
                'theoretical_qty': 0.0,
                'items': {},
                'lines': [],
                'closing': [],
            'ingredient_ids': set(),
            'zero_unit_cost': 0,
                'actual': 0.0,
                'daily': 0.0,
                'theo': 0.0,
                'purchase_qty': 0.0,
                'daily_qty': 0.0,
            }
            company_row['categories'][key] = row
            return row

        for company in selected:
            company_bucket(company)

        lines = self.env['cpabooks.recipe.consumption'].search([
            ('company_id', 'in', company_ids),
            ('date', '>=', date_from),
            ('date', '<=', date_to),
        ])
        counted_days = set()
        for line in lines:
            if not line.company_id:
                continue
            cat_key = line.category or 'other'
            company_row = company_bucket(line.company_id)
            category_row = category_bucket(company_row, cat_key)
            if line.recipe_id:
                item_key = 'r%s' % line.recipe_id.id
                item_name = line.recipe_id.name
                seen_recipes.add(line.recipe_id.id)
            else:
                item_key = 'p%s' % (line.sale_product_id.id or 0)
                item_name = line.sale_product_id.display_name or 'Sale with no recipe'
            item = category_row['items'].get(item_key)
            if not item:
                item = {
                    'key': item_key,
                    'name': item_name,
                    'recipe_id': line.recipe_id.id or 0,
                    'product_id': line.sale_product_id.id or 0,
                    'sold_qty': 0.0,
                    'sales': 0.0,
                    'cost': 0.0,
                    'theoretical_qty': 0.0,
                    'actual': 0.0,
                    'daily': 0.0,
                    'theo': 0.0,
                    'theory_by_product': {},
                    'gap_extra': 'No recipe' if not line.recipe_id else '',
                }
                category_row['items'][item_key] = item
            day_key = (item_key, line.date)
            if line.recipe_id:
                if day_key not in counted_days:
                    counted_days.add(day_key)
                    sales = line.sales_value or 0.0
                    sold = line.sold_qty or 0.0
                    item['sales'] += sales
                    item['sold_qty'] += sold
                    category_row['sales'] += sales
                    category_row['sold_qty'] += sold
                    company_row['sales'] += sales
                    company_row['sold_qty'] += sold
            else:
                sales = line.sales_value or 0.0
                sold = line.sold_qty or 0.0
                item['sales'] += sales
                item['sold_qty'] += sold
                category_row['sales'] += sales
                category_row['sold_qty'] += sold
                company_row['sales'] += sales
                company_row['sold_qty'] += sold
            cost = line.theoretical_cost or 0.0
            theory = line.theoretical_qty or 0.0
            company_row['theoretical_qty'] += theory
            item['cost'] += cost
            item['theoretical_qty'] += theory
            category_row['cost'] += cost
            category_row['theoretical_qty'] += theory
            company_row['cost'] += cost
            if line.ingredient_id:
                category_row['ingredient_ids'].add(line.ingredient_id.id)
                item.setdefault('theory_by_product', {})
                item['theory_by_product'][line.ingredient_id.id] = item['theory_by_product'].get(line.ingredient_id.id, 0.0) + theory
            notes = []
            if not line.sales_value:
                notes.append('Sold amount missing')
            if not line.theoretical_cost:
                notes.append('Cost missing')
            if line.ingredient_id and not line.unit_cost:
                notes.append('Unit cost missing')
                category_row['zero_unit_cost'] += 1
            if not line.ingredient_id:
                notes.append('No recipe')
            category_row['lines'].append({
                'date': line.date and line.date.strftime('%d/%m/%Y') or '',
                'sale': line.sale_product_id.display_name or '',
                'recipe': line.recipe_id.name or '',
                'ingredient': line.ingredient_id.display_name or 'No ingredient',
                'sold': qty(line.sold_qty),
                'recipe_qty': qty(line.recipe_qty),
                'theoretical_qty': qty(line.theoretical_qty),
                'unit_cost': qty(line.unit_cost),
                'cost': qty(line.theoretical_cost),
                'sales': qty(line.sales_value),
                'potential': '%.1f%%' % (line.potential_cost_pct or 0.0),
                'origin': line.origin or '',
                'gap': ', '.join(notes),
                'recipe_id': line.recipe_id.id or 0,
                'product_id': line.ingredient_id.id or 0,
                'theory_raw': theory,
                'potential_raw': cost,
                'sales_missing': not line.sales_value,
                'cost_missing': not line.theoretical_cost,
            })

        closings = self.env['cpabooks.recipe.closing'].search([
            ('company_id', 'in', company_ids),
            ('date', '>=', date_from),
            ('date', '<=', date_to),
        ])
        for closing in closings:
            if not closing.company_id:
                continue
            company_row = company_bucket(closing.company_id)
            payload = {
                'id': closing.id,
                'date': closing.date and closing.date.strftime('%d/%m/%Y') or '',
                'product': closing.product_id.display_name or '',
                'product_id': closing.product_id.id,
                'purchase': qty(closing.purchase_qty),
                'theoretical': qty(closing.theoretical_consumption),
                'closing': qty(closing.actual_closing_qty),
                'variance': qty(closing.variance_qty),
                'state': closing.state or '',
                'gap': '' if closing.theoretical_consumption else 'No expected use from sales',
                'purchase_missing': not closing.purchase_qty,
                'use_missing': not closing.theoretical_consumption,
            }
            attached = False
            for category_row in company_row['categories'].values():
                if closing.product_id.id in category_row['ingredient_ids']:
                    category_row['closing'].append(payload)
                    attached = True
            if attached:
                continue
            category_row = category_bucket(company_row, 'unmatched')
            category_row['closing'].append(payload)
            item_key = 'c%s' % closing.product_id.id
            if item_key not in category_row['items']:
                category_row['items'][item_key] = {
                    'key': item_key,
                    'name': closing.product_id.display_name or '',
                    'recipe_id': 0,
                    'product_id': closing.product_id.id or 0,
                    'sold_qty': 0.0,
                    'sales': 0.0,
                    'cost': 0.0,
                    'theoretical_qty': 0.0,
                    'actual': 0.0,
                    'daily': 0.0,
                    'theo': 0.0,
                    'theory_by_product': {},
                    'gap_extra': 'Closed or purchased, not on a sold recipe',
                }

        def fill_com():
            cost_type = self.env.ref('account.data_account_type_direct_costs', raise_if_not_found=False)
            if not cost_type:
                cost_type = self.env['account.account.type'].sudo().search([('name', '=', 'Cost of Revenue')], limit=1)
            domain = [
                ('move_id.state', '=', 'posted'),
                ('move_id.move_type', '=', 'in_invoice'),
                ('display_type', '=', False),
                ('product_id', '!=', False),
                ('company_id', 'in', company_ids or [0]),
                ('move_id.invoice_date', '>=', date_from),
                ('move_id.invoice_date', '<=', date_to),
            ]
            if cost_type:
                domain.append(('account_id.user_type_id', '=', cost_type.id))
            purchases = {}
            for row in self.env['account.move.line'].sudo().read_group(
                domain, ['price_subtotal:sum', 'quantity:sum'], ['company_id', 'product_id'], lazy=False,
            ):
                company = row.get('company_id') and row['company_id'][0]
                product = row.get('product_id') and row['product_id'][0]
                if not company or not product:
                    continue
                purchases[(company, product)] = {
                    'amount': row.get('price_subtotal') or 0.0,
                    'qty': row.get('quantity') or 0.0,
                }
            from dateutil.relativedelta import relativedelta
            month_from = date_to - relativedelta(months=1)
            month_domain = [
                ('move_id.state', '=', 'posted'),
                ('move_id.move_type', '=', 'in_invoice'),
                ('display_type', '=', False),
                ('product_id', '!=', False),
                ('company_id', 'in', company_ids or [0]),
                ('move_id.invoice_date', '>=', month_from),
                ('move_id.invoice_date', '<=', date_to),
            ]
            if cost_type:
                month_domain.append(('account_id.user_type_id', '=', cost_type.id))
            month_rates = {}
            for row in self.env['account.move.line'].sudo().read_group(
                month_domain, ['price_subtotal:sum', 'quantity:sum'], ['company_id', 'product_id'], lazy=False,
            ):
                company = row.get('company_id') and row['company_id'][0]
                product = row.get('product_id') and row['product_id'][0]
                if not company or not product:
                    continue
                month_rates[(company, product)] = {
                    'amount': row.get('price_subtotal') or 0.0,
                    'qty': row.get('quantity') or 0.0,
                }
            outs = {}
            daily_days = {}
            if 'daily.goods.consumption' in self.env:
                for rec in self.env['daily.goods.consumption'].sudo().search([
                    ('company_id', 'in', company_ids or [0]),
                    ('consume_date', '>=', date_from),
                    ('consume_date', '<=', date_to),
                ]):
                    key = (rec.company_id.id, rec.product_id.id)
                    outs[key] = outs.get(key, 0.0) + (rec.qty or 0.0)
                    daily_days.setdefault(key, []).append({
                        'sort': rec.consume_date,
                        'date': rec.consume_date and rec.consume_date.strftime('%d/%m/%Y') or '',
                        'qty': rec.qty or 0.0,
                    })

            def unit_of(cid, pid):
                info = purchases.get((cid, pid)) or {}
                if info.get('qty'):
                    return info['amount'] / info['qty'], 'period'
                info = month_rates.get((cid, pid)) or {}
                if info.get('qty'):
                    return info['amount'] / info['qty'], 'month'
                return 0.0, ''

            def touch(row):
                row.setdefault('actual', 0.0)
                row.setdefault('daily', 0.0)
                row.setdefault('theo', 0.0)

            for company_row in companies.values():
                cid = company_row['id']
                touch(company_row)
                slots = {}
                for cat in company_row['categories'].values():
                    touch(cat)
                    for item in cat['items'].values():
                        touch(item)
                        for pid, tqty in (item.get('theory_by_product') or {}).items():
                            slots.setdefault(pid, []).append((cat, item, tqty))
                product_total_theory = {pid: sum(slot[2] for slot in slot_list) for pid, slot_list in slots.items()}
                seen = set()

                def add_qty(row, purchase_qty, out_qty, share=1.0):
                    row['purchase_qty'] = row.get('purchase_qty', 0.0) + (purchase_qty or 0.0) * share
                    row['daily_qty'] = row.get('daily_qty', 0.0) + (out_qty or 0.0) * share

                def add_product(pid, amount, out_qty, purchase_qty):
                    unit, _rate_kind = unit_of(cid, pid)
                    daily_value = (out_qty or 0.0) * unit
                    company_row['actual'] += amount or 0.0
                    company_row['daily'] += daily_value
                    company_row['purchase_qty'] += purchase_qty or 0.0
                    company_row['daily_qty'] += out_qty or 0.0
                    slot_list = slots.get(pid) or []
                    total_t = product_total_theory.get(pid) or 0.0
                    if not slot_list:
                        product = self.env['product.product'].browse(pid)
                        cat = category_bucket(company_row, 'unmatched')
                        touch(cat)
                        item_key = 'b%s' % pid
                        item = cat['items'].get(item_key)
                        if not item:
                            item = {
                                'key': item_key,
                                'name': product.display_name or '',
                                'recipe_id': 0,
                                'product_id': pid,
                                'sold_qty': 0.0,
                                'sales': 0.0,
                                'cost': 0.0,
                                'theoretical_qty': 0.0,
                                'actual': 0.0,
                                'daily': 0.0,
                                'theo': 0.0,
                                'purchase_qty': 0.0,
                                'daily_qty': 0.0,
                                'theory_by_product': {},
                                'gap_extra': 'On Daily Goods, not on a sold recipe',
                            }
                            cat['items'][item_key] = item
                        item['actual'] += amount or 0.0
                        item['daily'] += daily_value
                        add_qty(item, purchase_qty, out_qty)
                        cat['actual'] += amount or 0.0
                        cat['daily'] += daily_value
                        add_qty(cat, purchase_qty, out_qty)
                        return
                    for cat, item, tqty in slot_list:
                        share = (tqty / total_t) if total_t else (1.0 / len(slot_list))
                        item['actual'] += (amount or 0.0) * share
                        item['daily'] += daily_value * share
                        item['theo'] += tqty * unit
                        add_qty(item, purchase_qty, out_qty, share)
                        cat['actual'] += (amount or 0.0) * share
                        cat['daily'] += daily_value * share
                        cat['theo'] += tqty * unit
                        add_qty(cat, purchase_qty, out_qty, share)
                        company_row['theo'] += tqty * unit

                for (pcid, pid), info in purchases.items():
                    if pcid != cid:
                        continue
                    seen.add(pid)
                    add_product(pid, info['amount'], outs.get((cid, pid), 0.0), info.get('qty') or 0.0)
                for (pcid, pid), out_qty in outs.items():
                    if pcid != cid or pid in seen:
                        continue
                    add_product(pid, 0.0, out_qty, 0.0)
                name_of = {}

                def product_name(ipid):
                    if ipid not in name_of:
                        name_of[ipid] = self.env['product.product'].browse(ipid).display_name or ''
                    return name_of[ipid]

                for cat in company_row['categories'].values():
                    for item in cat['items'].values():
                        pid = item.get('product_id') or 0
                        theory = item.get('theory_by_product') or {}
                        if theory:
                            ingredient_ids = list(theory.keys())
                        elif pid and not product_total_theory.get(pid):
                            ingredient_ids = [pid]
                        elif pid and (item.get('daily') or item.get('purchase_qty')):
                            ingredient_ids = [pid]
                        else:
                            ingredient_ids = []
                        detail_lines = []
                        for ipid in ingredient_ids:
                            days = daily_days.get((cid, ipid)) or []
                            unit, rate_kind = unit_of(cid, ipid)
                            if rate_kind == 'month':
                                rate_note = '1-month average rate'
                            elif rate_kind == 'period':
                                rate_note = ''
                            else:
                                rate_note = 'No purchase in this period or the last month'
                            total_t = product_total_theory.get(ipid) or 0.0
                            own = theory.get(ipid) or 0.0
                            share = (own / total_t) if total_t else 1.0
                            for day in sorted(days, key=lambda one: one.get('sort') or ''):
                                day_qty = (day.get('qty') or 0.0) * share
                                detail_lines.append({
                                    'date': day['date'],
                                    'name': product_name(ipid),
                                    'product_id': ipid,
                                    'qty': qty(day_qty),
                                    'unit': qty(unit),
                                    'amount': money(day_qty * unit),
                                    'note': rate_note,
                                    'rate': rate_kind,
                                })
                        item['daily_lines'] = detail_lines
                for cat in list(company_row['categories'].values()):
                    for line in cat['lines']:
                        pid = line.get('product_id') or 0
                        total_t = product_total_theory.get(pid) or 0.0
                        info = purchases.get((cid, pid)) or {}
                        amount = info.get('amount') or 0.0
                        unit, _rate_kind = unit_of(cid, pid)
                        share = (line.get('theory_raw') or 0.0) / total_t if total_t else 0.0
                        line['actual'] = money(amount * share)
                        line['potential_com'] = money(line.get('potential_raw') or 0.0)
                        line['daily'] = money((outs.get((cid, pid), 0.0) * unit) * share)
                        line['daily_qty'] = qty((outs.get((cid, pid), 0.0) or 0.0) * share)
                        line['theo'] = money((line.get('theory_raw') or 0.0) * unit)
                        line['actual_missing'] = not (amount * share)
                        line['potential_missing'] = not line.get('potential_raw')
                        line['daily_missing'] = not ((outs.get((cid, pid), 0.0) * unit) * share)
                        line['theo_missing'] = not ((line.get('theory_raw') or 0.0) * unit)

        fill_com()

        def com_gap(sales, actual, potential, daily, theo):
            if not sales and not actual and not potential and not daily and not theo:
                return 'No sales and no cost'
            notes = []
            if sales and not actual:
                notes.append('Actual purchase missing')
            if sales and not potential:
                notes.append('Potential missing')
            if (actual or potential) and not daily:
                notes.append('Daily consumption not entered')
            if potential and not theo:
                notes.append('Theoretical missing')
            return ', '.join(notes)

        def pack_item(item):
            gap = com_gap(item['sales'], item.get('actual'), item['cost'], item.get('daily'), item.get('theo'))
            extra = item.get('gap_extra') or ''
            if any(day.get('rate') == 'month' for day in item.get('daily_lines') or []):
                extra = (extra + ', 1-month average rate').strip(', ')
            elif item.get('daily_lines') and not item.get('daily'):
                extra = (extra + ', Daily qty entered, no rate in the last month').strip(', ')
            if extra and extra not in gap:
                gap = (gap + ', ' + extra).strip(', ')
            ingredient_ids = list((item.get('theory_by_product') or {}).keys())
            if item.get('product_id') and item.get('product_id') not in ingredient_ids:
                ingredient_ids.append(item.get('product_id'))
            return {
                'key': item['key'],
                'name': item['name'],
                'recipe_id': item['recipe_id'],
                'product_id': item.get('product_id') or 0,
                'drill_products': ingredient_ids,
                'sold_qty': qty_show(item['sold_qty']),
                'actual_qty': qty_show(item.get('purchase_qty')),
                'daily_qty': qty_show(item.get('daily_qty')),
                'theo_qty': qty_show(item['theoretical_qty']),
                'sales': money(item['sales']),
                'actual': money(item.get('actual')),
                'potential': money(item['cost']),
                'daily': money(item.get('daily')),
                'theo': money(item.get('theo')),
                'cost': money(item['cost']),
                'theoretical_qty': qty_show(item['theoretical_qty']),
                'potential_pct': '%.1f%%' % percent(item['cost'], item['sales']),
                'com_pct': '%.1f%%' % percent(item.get('actual') or 0.0, item['sales']),
                'gm': money((item['sales'] or 0.0) - (item.get('actual') or 0.0)),
                'gap': gap,
                'sales_missing': not item['sales'],
                'actual_missing': not item.get('actual'),
                'potential_missing': not item['cost'],
                'daily_missing': not item.get('daily'),
                'theo_missing': not item.get('theo'),
                'cost_missing': not item['cost'],
                'daily_lines': item.get('daily_lines') or [],
            }

        def pack_category(category_row):
            gap = com_gap(category_row['sales'], category_row.get('actual'), category_row['cost'], category_row.get('daily'), category_row.get('theo'))
            if category_row['zero_unit_cost']:
                extra = '%s lines with no unit cost' % category_row['zero_unit_cost']
                gap = (gap + ', ' + extra).strip(', ')
            if category_row['key'] == 'unmatched' and (category_row['closing'] or category_row.get('actual') or category_row.get('daily')):
                gap = (gap + ', On Daily Goods, not on a sold recipe').strip(', ')
            items = [pack_item(item) for item in category_row['items'].values()]
            items.sort(key=lambda row: row['name'])
            drill_products = []
            for packed_item in items:
                for product_id in packed_item.get('drill_products') or []:
                    if product_id not in drill_products:
                        drill_products.append(product_id)
            return {
                'key': category_row['key'],
                'name': category_row['name'],
                'drill_products': drill_products,
                'sold_qty': qty_show(category_row['sold_qty']),
                'actual_qty': qty_show(category_row.get('purchase_qty')),
                'daily_qty': qty_show(category_row.get('daily_qty')),
                'theo_qty': qty_show(category_row['theoretical_qty']),
                'sales': money(category_row['sales']),
                'actual': money(category_row.get('actual')),
                'potential': money(category_row['cost']),
                'daily': money(category_row.get('daily')),
                'theo': money(category_row.get('theo')),
                'cost': money(category_row['cost']),
                'theoretical_qty': qty_show(category_row['theoretical_qty']),
                'potential_pct': '%.1f%%' % percent(category_row['cost'], category_row['sales']),
                'com_pct': '%.1f%%' % percent(category_row.get('actual') or 0.0, category_row['sales']),
                'gm': money((category_row['sales'] or 0.0) - (category_row.get('actual') or 0.0)),
                'gap': gap,
                'sales_missing': not category_row['sales'],
                'actual_missing': not category_row.get('actual'),
                'potential_missing': not category_row['cost'],
                'daily_missing': not category_row.get('daily'),
                'theo_missing': not category_row.get('theo'),
                'cost_missing': not category_row['cost'],
                'items': items,
                'lines': category_row['lines'],
                'closing': category_row['closing'],
            }

        packed = []
        total_sales = 0.0
        total_actual = 0.0
        total_potential = 0.0
        total_daily = 0.0
        total_theo = 0.0
        total_sold_qty = 0.0
        total_purchase_qty = 0.0
        total_daily_qty = 0.0
        total_theo_qty = 0.0
        for company_row in companies.values():
            categories = [pack_category(row) for row in company_row['categories'].values()]
            categories.sort(key=lambda row: row['name'])
            actual = company_row.get('actual') or 0.0
            potential = company_row['cost']
            daily = company_row.get('daily') or 0.0
            theo = company_row.get('theo') or 0.0
            total_sales += company_row['sales']
            total_actual += actual
            total_potential += potential
            total_daily += daily
            total_theo += theo
            total_sold_qty += company_row.get('sold_qty') or 0.0
            total_purchase_qty += company_row.get('purchase_qty') or 0.0
            total_daily_qty += company_row.get('daily_qty') or 0.0
            total_theo_qty += company_row.get('theoretical_qty') or 0.0
            packed.append({
                'id': company_row['id'],
                'name': company_row['name'],
                'drill_products': [],
                'sold_qty': qty_show(company_row.get('sold_qty')),
                'actual_qty': qty_show(company_row.get('purchase_qty')),
                'daily_qty': qty_show(company_row.get('daily_qty')),
                'theo_qty': qty_show(company_row.get('theoretical_qty')),
                'sales': money(company_row['sales']),
                'actual': money(actual),
                'potential': money(potential),
                'daily': money(daily),
                'theo': money(theo),
                'cost': money(actual),
                'com_pct': '%.1f%%' % percent(actual, company_row['sales']),
                'gm': money(company_row['sales'] - actual),
                'gap': com_gap(company_row['sales'], actual, potential, daily, theo),
                'sales_missing': not company_row['sales'],
                'actual_missing': not actual,
                'potential_missing': not potential,
                'daily_missing': not daily,
                'theo_missing': not theo,
                'cost_missing': not actual,
                'categories': categories,
            })
        packed.sort(key=lambda row: row['name'])
        return {
            'date_from': fields.Date.to_string(date_from),
            'date_to': fields.Date.to_string(date_to),
            'companies': packed,
            'total': {
                'sales': money(total_sales),
                'actual': money(total_actual),
                'potential': money(total_potential),
                'daily': money(total_daily),
                'theo': money(total_theo),
                'sold_qty': qty_show(total_sold_qty),
                'actual_qty': qty_show(total_purchase_qty),
                'daily_qty': qty_show(total_daily_qty),
                'theo_qty': qty_show(total_theo_qty),
                'cost': money(total_actual),
                'com_pct': '%.1f%%' % percent(total_actual, total_sales),
                'gm': money(total_sales - total_actual),
                'gap': com_gap(total_sales, total_actual, total_potential, total_daily, total_theo),
            },
        }

    @api.model
    def get_key_result_purchases(self, company_id, product_ids=None, date_from=None, date_to=None, kind='actual'):
        today = fields.Date.context_today(self)
        date_from = fields.Date.to_date(date_from) if date_from else today.replace(day=1)
        date_to = fields.Date.to_date(date_to) if date_to else today
        company_id = int(company_id or 0)
        if company_id not in self.env.companies.ids:
            company_id = self.env.companies[:1].id
        product_ids = [int(pid) for pid in (product_ids or []) if int(pid or 0)]
        cost_type = self.env.ref('account.data_account_type_direct_costs', raise_if_not_found=False)
        if not cost_type:
            cost_type = self.env['account.account.type'].sudo().search([('name', '=', 'Cost of Revenue')], limit=1)

        def bill_domain(start, end, products):
            domain = [
                ('move_id.state', '=', 'posted'),
                ('move_id.move_type', '=', 'in_invoice'),
                ('display_type', '=', False),
                ('product_id', '!=', False),
                ('company_id', '=', company_id),
                ('move_id.invoice_date', '>=', start),
                ('move_id.invoice_date', '<=', end),
            ]
            if products:
                domain.append(('product_id', 'in', products))
            if cost_type:
                domain.append(('account_id.user_type_id', '=', cost_type.id))
            return domain

        windows = [(date_from, date_to, product_ids)]
        window_note = ''
        if kind != 'actual' and product_ids:
            from dateutil.relativedelta import relativedelta
            month_from = date_to - relativedelta(months=1)
            have_qty = set()
            for row in self.env['account.move.line'].sudo().read_group(
                bill_domain(date_from, date_to, product_ids), ['quantity:sum'], ['product_id'], lazy=False,
            ):
                product = row.get('product_id') and row['product_id'][0]
                if product and row.get('quantity'):
                    have_qty.add(product)
            period_products = [pid for pid in product_ids if pid in have_qty]
            month_products = [pid for pid in product_ids if pid not in have_qty]
            windows = []
            if period_products:
                windows.append((date_from, date_to, period_products))
            if month_products:
                start = month_from if month_from < date_from else date_from
                windows.append((start, date_to, month_products))
                if month_from < date_from:
                    window_note = 'Where the period has no bill, the list uses that item\'s last 1 month average.'
            if not windows:
                windows = [(date_from, date_to, product_ids)]
        records = self.env['account.move.line']
        seen_ids = set()
        window_labels = []
        for start, end, products in windows:
            window_labels.append('%s to %s' % (start.strftime('%d/%m/%Y'), end.strftime('%d/%m/%Y')))
            found = self.env['account.move.line'].sudo().search(
                bill_domain(start, end, products), order='move_id.invoice_date, move_id.name, id',
            )
            extra = found.filtered(lambda line: line.id not in seen_ids)
            seen_ids.update(extra.ids)
            records |= extra
        records = records.sorted(key=lambda line: (line.move_id.invoice_date or date_from, line.move_id.name or '', line.id))
        lines = []
        total_qty = 0.0
        total_amount = 0.0
        for line in records:
            qty_value = line.quantity or 0.0
            amount_value = line.price_subtotal or 0.0
            total_qty += qty_value
            total_amount += amount_value
            bill = line.move_id
            lines.append({
                'date': bill.invoice_date and bill.invoice_date.strftime('%d/%m/%Y') or '',
                'bill': bill.name or bill.ref or '',
                'partner': bill.partner_id.display_name or '',
                'product': line.product_id.display_name or '',
                'qty': '{:,.2f}'.format(qty_value),
                'amount': '{:,.2f}'.format(amount_value),
            })
        title = 'Purchases'
        if len(product_ids) == 1:
            title = self.env['product.product'].browse(product_ids[0]).display_name or title
        return {
            'title': title,
            'window': ' · '.join(window_labels),
            'window_note': window_note,
            'lines': lines,
            'total_qty': '{:,.2f}'.format(total_qty),
            'total_amount': '{:,.2f}'.format(total_amount),
        }

    @api.model
    def get_product_card(self, product_id, company_id=None):
        product = self.env['product.product'].browse(int(product_id or 0)).exists()
        if not product:
            return {}
        product = product.sudo()
        if company_id:
            product = product.with_context(force_company=int(company_id))
        kind = {'consu': 'Consumable', 'service': 'Service', 'product': 'Storable'}.get(product.type, product.type or '')
        return {
            'id': product.id,
            'tmpl_id': product.product_tmpl_id.id,
            'name': product.display_name,
            'code': product.default_code or '',
            'barcode': product.barcode or '',
            'category': product.categ_id.display_name or '',
            'uom': product.uom_id.name or '',
            'sale_price': '{:,.2f}'.format(product.lst_price or 0.0),
            'cost': '{:,.2f}'.format(product.standard_price or 0.0),
            'qty': '{:,.2f}'.format(product.qty_available or 0.0),
            'kind': kind,
            'sale_ok': bool(product.sale_ok),
            'purchase_ok': bool(product.purchase_ok),
        }

    @api.model
    def export_key_results_xlsx(self, date_from=None, date_to=None):
        import base64
        import io

        from openpyxl import Workbook
        from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
        from openpyxl.utils import get_column_letter

        data = self.get_key_results(date_from, date_to)
        book = Workbook()
        sheet = book.active
        sheet.title = 'Key Results'
        thin = Border(
            left=Side(style='thin', color='D5DEE3'),
            right=Side(style='thin', color='D5DEE3'),
            top=Side(style='thin', color='D5DEE3'),
            bottom=Side(style='thin', color='D5DEE3'),
        )
        fill_name = PatternFill('solid', fgColor='1A7F86')
        fill_sale = PatternFill('solid', fgColor='1565C0')
        fill_cost = PatternFill('solid', fgColor='EF6C00')
        fill_group = PatternFill('solid', fgColor='FFF2CC')
        font_head = Font(bold=True, color='FFFFFF')
        font_group = Font(bold=True, color='000000')
        font_item = Font(color='000000')
        align_left = Alignment(horizontal='left', vertical='center')
        align_right = Alignment(horizontal='right', vertical='center')

        def amount(value):
            text = str(value or '0').replace(',', '').replace('%', '').strip()
            try:
                return float(text or 0)
            except ValueError:
                return 0.0

        def paint(cell, kind, group):
            cell.border = thin
            cell.alignment = align_right if kind in ('sale', 'cost', 'pct') else align_left
            if group == 'head':
                cell.font = font_head
                cell.fill = {'sale': fill_sale, 'cost': fill_cost, 'pct': fill_cost}.get(kind, fill_name)
            elif group:
                cell.font = font_group
                cell.fill = fill_group
            else:
                cell.font = font_item

        def write_row(row_no, values, kinds, group):
            for col, value in enumerate(values, 1):
                cell = sheet.cell(row_no, col, value)
                kind = kinds[col - 1]
                if kind in ('sale', 'cost') and group != 'head':
                    cell.value = amount(value)
                    cell.number_format = '#,##0.00'
                elif kind == 'pct' and group != 'head':
                    cell.value = amount(value)
                    cell.number_format = '0.0"%"'
                paint(cell, kind, group)

        metric_head = [
            'Level', 'Company', 'Category', 'Item', 'Total Sales',
            'COM (Actual)', 'COM Potential', 'COM (Daily Consumptions)',
            'COM Theoretical', 'COM %', 'GM', 'Missing',
        ]
        metric_kinds = [
            'text', 'text', 'text', 'text', 'sale',
            'cost', 'cost', 'cost', 'cost', 'pct', 'sale', 'text',
        ]
        row_no = 1
        write_row(row_no, metric_head, metric_kinds, 'head')
        for company in data.get('companies') or []:
            row_no += 1
            write_row(row_no, [
                'Company', company.get('name'), '', '', company.get('sales'),
                company.get('actual'), company.get('potential'), company.get('daily'),
                company.get('theo'), company.get('com_pct'), company.get('gm'), company.get('gap'),
            ], metric_kinds, 'group')
            for category in company.get('categories') or []:
                row_no += 1
                write_row(row_no, [
                    'Category', company.get('name'), category.get('name'), '', category.get('sales'),
                    category.get('actual'), category.get('potential'), category.get('daily'),
                    category.get('theo'), category.get('com_pct'), category.get('gm'), category.get('gap'),
                ], metric_kinds, 'group')
                for item in category.get('items') or []:
                    row_no += 1
                    write_row(row_no, [
                        'Item', company.get('name'), category.get('name'), item.get('name'), item.get('sales'),
                        item.get('actual'), item.get('potential'), item.get('daily'),
                        item.get('theo'), item.get('com_pct'), item.get('gm'), item.get('gap'),
                    ], metric_kinds, False)
        total = data.get('total') or {}
        row_no += 1
        write_row(row_no, [
            'Total', '', '', '', total.get('sales'), total.get('actual'), total.get('potential'),
            total.get('daily'), total.get('theo'), total.get('com_pct'), total.get('gm'), total.get('gap'),
        ], metric_kinds, 'group')

        row_no += 2
        qty_head = [
            'Company', 'Category', 'Item', 'Date', 'Ingredient', 'Sold qty', 'Theoretical qty',
            'COM (Actual)', 'COM Potential', 'COM (Daily Consumptions)', 'COM Theoretical', 'Missing',
        ]
        qty_kinds = [
            'text', 'text', 'text', 'text', 'text', 'sale', 'cost',
            'cost', 'cost', 'cost', 'cost', 'text',
        ]
        write_row(row_no, qty_head, qty_kinds, 'head')
        for company in data.get('companies') or []:
            for category in company.get('categories') or []:
                lines = category.get('lines') or []
                if not lines:
                    continue
                row_no += 1
                write_row(row_no, [
                    'Category', company.get('name'), category.get('name'), '', '',
                    '', '', '', '', '', '', '',
                ], qty_kinds, 'group')
                for line in lines:
                    row_no += 1
                    gap = ((line.get('origin') + '. ') if line.get('origin') else '') + (line.get('gap') or '')
                    write_row(row_no, [
                        company.get('name'), category.get('name'), line.get('sale') or line.get('recipe') or '',
                        line.get('date'), line.get('ingredient'), line.get('sold'), line.get('theoretical_qty'),
                        line.get('actual'), line.get('potential_com'), line.get('daily'), line.get('theo'), gap,
                    ], qty_kinds, False)
        row_no += 2
        close_head = [
            'Company', 'Category', 'Date', 'Ingredient', 'Purchase qty', 'Expected qty',
            'Chef closing qty', 'Variance', 'Status', 'Missing',
        ]
        close_kinds = [
            'text', 'text', 'text', 'text', 'cost', 'cost', 'cost', 'cost', 'text', 'text',
        ]
        write_row(row_no, close_head, close_kinds, 'head')
        for company in data.get('companies') or []:
            for category in company.get('categories') or []:
                closings = category.get('closing') or []
                if not closings:
                    continue
                row_no += 1
                write_row(row_no, [
                    'Category', company.get('name'), category.get('name'), '', '', '', '', '', '', '',
                ], close_kinds, 'group')
                for line in closings:
                    row_no += 1
                    write_row(row_no, [
                        company.get('name'), category.get('name'), line.get('date'), line.get('product'),
                        line.get('purchase'), line.get('theoretical'), line.get('closing'),
                        line.get('variance'), line.get('state'), line.get('gap'),
                    ], close_kinds, False)

        widths = [16, 42, 24, 32, 16, 18, 18, 28, 18, 12, 16, 42]
        for index, width in enumerate(widths, 1):
            sheet.column_dimensions[get_column_letter(index)].width = width
        sheet.auto_filter.ref = 'A1:L1'
        sheet.freeze_panes = 'A2'
        sheet.sheet_view.showGridLines = False
        buffer = io.BytesIO()
        book.save(buffer)
        filename = 'key_results_%s_%s.xlsx' % (data.get('date_from') or '', data.get('date_to') or '')
        return {'filename': filename, 'file': base64.b64encode(buffer.getvalue()).decode('ascii')}
